"""文件測試用的自造樣本（全部在測試中程式產生，不讀任何真實文件）。

與 `tests/documents/conftest.py` 的 fixture 同一套做法；這裡是給套件內
（`tests/api`、`tests/mcp`）以相對 import 共用的純函式版本。
"""

from __future__ import annotations

import io


def pdf_bytes(pages: list[str]) -> bytes:
    """以 Helvetica 寫 ASCII 內容流的 pdf（空字串 = 空白頁）。"""
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, StreamObject

    writer = PdfWriter()
    font = writer._add_object(
        DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
    )
    for text in pages:
        page = writer.add_blank_page(612, 792)
        if not text:
            continue
        stream = StreamObject()
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1"))
        page[NameObject("/Contents")] = writer._add_object(stream)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def docx_bytes(sections: list[tuple[str, str]]) -> bytes:
    """[(Heading 1 標題, 內文段落)]。"""
    import docx

    document = docx.Document()
    for heading, body in sections:
        document.add_heading(heading, level=1)
        document.add_paragraph(body)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def pptx_bytes(slides: list[tuple[str, str]]) -> bytes:
    """[(投影片標題, 內文)]。"""
    import pptx

    presentation = pptx.Presentation()
    for title, body in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    buf = io.BytesIO()
    presentation.save(buf)
    return buf.getvalue()
