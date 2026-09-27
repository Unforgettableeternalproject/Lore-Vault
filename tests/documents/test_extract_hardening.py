"""抽取器的惡意輸入防護：壓縮炸彈（竄改宣告大小）、XXE、billion laughs。

樣本全部在測試中程式產生。記憶體以 tracemalloc 量測峰值（Python 物件配置，
解壓出的 bytes 會算在內）。
"""

from __future__ import annotations

import io
import struct
import time
import tracemalloc
import zipfile

import pytest

from lore_vault.documents.extract import (
    CORRUPT,
    TOO_LARGE,
    ExtractionError,
    Limits,
    extract,
)
from lore_vault.documents.extract.base import check_ooxml_container

BOMB_SIZE = 50 * 1024 * 1024
# 修正後的峰值應在串流分塊量級；修正前約為 BOMB_SIZE
PEAK_LIMIT = 8 * 1024 * 1024


def _lying_zip(member: str, size: int, declared: int) -> bytes:
    """單一成員的 DEFLATE zip，local header 與 central directory 的
    uncompressed size 都竄改成 `declared`（CRC 仍是真實內容的）。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(member, bytes(size))
    data = bytearray(buf.getvalue())
    assert data[:4] == b"PK\x03\x04"
    struct.pack_into("<I", data, 22, declared)
    central = data.index(b"PK\x01\x02")
    struct.pack_into("<I", data, central + 24, declared)
    return bytes(data)


@pytest.fixture(scope="module")
def lying_bomb() -> bytes:
    # python-docx 開檔第一個讀的就是 [Content_Types].xml
    data = _lying_zip("[Content_Types].xml", BOMB_SIZE, 100)
    assert len(data) < 100_000
    return data


def _peak_of(func):
    # 先載入解析套件，首次 import 的配置不算進峰值
    import docx  # noqa: F401
    import lxml.etree  # noqa: F401
    import pptx  # noqa: F401

    tracemalloc.start()
    tracemalloc.reset_peak()
    try:
        try:
            outcome = func()
        except ExtractionError as exc:
            outcome = exc
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    return outcome, peak


@pytest.mark.parametrize("name", ["bomb.docx", "bomb.pptx"])
def test_lying_zip_bomb_rejected_with_bounded_memory(lying_bomb, name):
    """宣告 100 位元組、實際 50MB：修正前檢查放行、python-docx 把 50MB 整個解進
    記憶體才因 CRC 失敗；修正後串流讀取，峰值在分塊量級。"""
    outcome, peak = _peak_of(lambda: extract(lying_bomb, name))
    assert isinstance(outcome, ExtractionError)
    assert outcome.code in (CORRUPT, TOO_LARGE)
    assert peak < PEAK_LIMIT, f"峰值 {peak} 位元組"


def test_declared_size_mismatch_is_corrupt(lying_bomb):
    with pytest.raises(ExtractionError) as info:
        check_ooxml_container(lying_bomb, Limits(), label="docx")
    assert info.value.code == CORRUPT


def test_declared_size_larger_than_actual_is_corrupt():
    data = _lying_zip("word/document.xml", 1000, 5000)
    with pytest.raises(ExtractionError) as info:
        check_ooxml_container(data, Limits(), label="docx")
    assert info.value.code == CORRUPT


def test_actual_unzipped_total_over_limit_is_too_large_and_streamed():
    """宣告值誠實但總量超限：宣告值就擋下（不解壓）。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
        for i in range(4):
            archive.writestr(f"part{i}.xml", bytes(300_000))
    limits = Limits(max_bytes=100_000, unzip_ratio=8)
    with pytest.raises(ExtractionError) as info:
        check_ooxml_container(buf.getvalue(), limits, label="docx")
    assert info.value.code == TOO_LARGE


def test_too_many_members_is_too_large():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        for i in range(21):
            archive.writestr(f"m{i}.xml", b"<a/>")
    with pytest.raises(ExtractionError) as info:
        check_ooxml_container(buf.getvalue(), Limits(max_zip_members=20), label="docx")
    assert info.value.code == TOO_LARGE
    check_ooxml_container(buf.getvalue(), Limits(max_zip_members=21), label="docx")


def test_member_compression_ratio_over_limit_is_too_large():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", bytes(4 * 1024 * 1024))
    data = buf.getvalue()
    with pytest.raises(ExtractionError) as info:
        check_ooxml_container(data, Limits(max_member_ratio=100), label="docx")
    assert info.value.code == TOO_LARGE
    # 上限放寬就通過（全零 DEFLATE 約 1000:1）
    check_ooxml_container(data, Limits(max_member_ratio=2000), label="docx")


def test_normal_office_files_pass_container_check(make_docx, make_pptx):
    docx_bytes = make_docx(lambda d: d.add_paragraph("正常內容" * 200))
    pptx_bytes = make_pptx(lambda p: p.slides.add_slide(p.slide_layouts[0]))
    check_ooxml_container(docx_bytes, Limits(), label="docx")
    check_ooxml_container(pptx_bytes, Limits(), label="pptx")


# ── XML 實體：XXE 與 billion laughs ─────────────────────────────────


def _rewrite_member(data: bytes, member: str, transform) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            content = src.read(info.filename)
            if info.filename == member:
                content = transform(content)
            dst.writestr(info.filename, content)
    return out.getvalue()


def _with_doctype(xml: bytes, root: str, subset: str, marker: bytes, ref: str):
    head, sep, rest = xml.partition(b"?>")
    assert sep, "缺 XML 宣告"
    doctype = f"\n<!DOCTYPE {root} [\n{subset}\n]>".encode()
    return (head + sep + doctype + rest).replace(marker, ref.encode())


_LOL = "\n".join(
    ['<!ENTITY lol0 "lol">']
    + [
        f'<!ENTITY lol{i} "' + f"&lol{i - 1};" * 10 + '">'  # 10^9 次展開
        for i in range(1, 10)
    ]
)


def _xxe_subset(secret_path) -> str:
    return f'<!ENTITY xxe SYSTEM "{secret_path.as_uri()}">'


def _office_with_entities(kind, make_docx, make_pptx, subset, ref):
    marker = "ENTITYMARKER"
    if kind == "docx":

        def build(document):
            document.add_paragraph("正常的第一段內容，確保文件不是空的")
            document.add_paragraph(marker)

        data, member, root = make_docx(build), "word/document.xml", "w:document"
    else:

        def build(presentation):
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "正常的投影片標題"
            slide.placeholders[1].text = marker

        data, member, root = make_pptx(build), "ppt/slides/slide1.xml", "p:sld"
    return _rewrite_member(
        data,
        member,
        lambda xml: _with_doctype(xml, root, subset, marker.encode(), ref),
    )


def _text_or_error(data: bytes, name: str) -> str:
    try:
        return "\n".join(s.text for s in extract(data, name).segments)
    except ExtractionError as exc:
        return f"<error {exc.code}>"


@pytest.mark.parametrize("kind", ["docx", "pptx"])
def test_external_entity_is_not_resolved(tmp_path, make_docx, make_pptx, kind):
    secret = tmp_path / "secret.txt"
    secret.write_text("TOP-SECRET-XXE-CONTENT", encoding="utf-8")
    data = _office_with_entities(
        kind, make_docx, make_pptx, _xxe_subset(secret), "&xxe;"
    )
    text = _text_or_error(data, f"xxe.{kind}")
    assert "TOP-SECRET-XXE-CONTENT" not in text


@pytest.mark.parametrize("kind", ["docx", "pptx"])
def test_billion_laughs_is_not_expanded(make_docx, make_pptx, kind):
    data = _office_with_entities(kind, make_docx, make_pptx, _LOL, "&lol9;")
    started = time.monotonic()
    outcome, peak = _peak_of(lambda: _text_or_error(data, f"lol.{kind}"))
    elapsed = time.monotonic() - started
    assert isinstance(outcome, str)
    assert "lollollol" not in outcome
    assert elapsed < 10
    assert peak < PEAK_LIMIT, f"峰值 {peak} 位元組"
