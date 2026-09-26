"""T-62：docx／pptx 抽取（自造樣本），含加密、損毀、副檔名不符與壓縮炸彈。"""

from __future__ import annotations

import io
import zipfile

import pytest
from pptx.util import Inches

from lore_vault.documents.extract import (
    CORRUPT,
    EMPTY_EXTRACTION,
    ENCRYPTED,
    TOO_LARGE,
    UNSUPPORTED_FORMAT,
    ExtractionError,
    Limits,
    extract,
)

OLE_MAGIC = bytes.fromhex("d0cf11e0a1b11ae1")


def _code(data: bytes, name: str, **limits) -> str:
    with pytest.raises(ExtractionError) as info:
        extract(data, name, limits=Limits(**limits) if limits else None)
    return info.value.code


def _locators(result):
    return [s.locator.to_dict() for s in result.segments]


# ── docx ────────────────────────────────────────────────────────────


def _structured_doc(document):
    document.add_paragraph("前言：這份文件說明記憶系統的部署方式與注意事項。")
    document.add_heading("第一章 系統設定", level=1)
    document.add_paragraph("資料庫使用 SQLite，並開啟 WAL 模式。")
    document.add_heading("資料庫", level=2)
    table = document.add_table(rows=2, cols=3)
    for col, text in enumerate(("欄位", "型別", "說明")):
        table.cell(0, col).text = text
    table.cell(1, 0).text = "id"
    table.cell(1, 1).text = "TEXT"
    table.cell(1, 1).merge(table.cell(1, 2))
    document.add_heading("第二章 部署", level=1)
    document.add_paragraph("使用 docker compose 啟動服務。")
    section = document.sections[0]
    section.header.paragraphs[0].text = "內部文件 頁首"
    section.footer.paragraphs[0].text = "頁尾：版權所有"


def test_docx_headings_tables_header_footer(make_docx):
    result = extract(make_docx(_structured_doc), "部署.docx")
    assert result.format == "docx"
    assert _locators(result) == [
        {"kind": "offset", "value": 0},
        {"kind": "heading", "value": "第一章 系統設定"},
        {"kind": "heading", "value": "第一章 系統設定 > 資料庫"},
        {"kind": "heading", "value": "第二章 部署"},
        {"kind": "header", "value": 1},
        {"kind": "footer", "value": 1},
    ]
    table_segment = result.segments[2].text
    assert table_segment.split("\n") == ["資料庫", "欄位 | 型別 | 說明", "id | TEXT"]
    assert result.segments[4].text == "內部文件 頁首"


def test_docx_without_headings_is_single_offset_segment(make_docx):
    def build(document):
        document.add_paragraph("第一段沒有任何標題樣式的內容，用來測試退回位移定位。")
        document.add_paragraph("第二段，仍然屬於同一個區塊，之後由切段步驟依長度處理。")

    result = extract(make_docx(build), "無標題.docx")
    assert _locators(result) == [{"kind": "offset", "value": 0}]
    assert result.segments[0].text.count("\n") == 1


def test_short_docx_is_not_blocked_by_min_chars(make_docx):
    """B1 裁決：min_chars 只套 pdf；只有一個字的 docx 仍是可檢索內容。"""
    result = extract(make_docx(lambda d: d.add_paragraph("圖")), "scan.docx")
    assert result.char_count == 1


def test_short_pptx_is_not_blocked_by_min_chars(make_pptx):
    def build(presentation):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "結論"

    result = extract(make_pptx(build), "短簡報.pptx")
    assert [s.locator.to_dict() for s in result.segments] == [
        {"kind": "slide", "value": 1}
    ]
    assert result.char_count == 2 and result.encoding is None


def test_docx_without_any_text_is_empty_extraction(make_docx):
    assert _code(make_docx(lambda d: d.add_paragraph("   ")), "blank.docx") == (
        EMPTY_EXTRACTION
    )


def test_docx_error_classification(make_docx, make_pptx):
    good = make_docx(_structured_doc)
    encrypted = OLE_MAGIC + b"\x00" * 64 + "EncryptedPackage".encode("utf-16-le")
    assert _code(encrypted, "secret.docx") == ENCRYPTED
    assert _code(OLE_MAGIC + b"\x00" * 64, "legacy.docx") == UNSUPPORTED_FORMAT
    assert _code(b"plain text pretending", "fake.docx") == UNSUPPORTED_FORMAT
    assert _code(good[: len(good) // 2], "broken.docx") == CORRUPT
    assert _code(b"PK\x03\x04" + b"\x00" * 40, "broken2.docx") == CORRUPT
    pptx_bytes = make_pptx(lambda p: p.slides.add_slide(p.slide_layouts[0]))
    assert _code(pptx_bytes, "renamed.docx") == UNSUPPORTED_FORMAT


def test_zip_without_word_parts_is_corrupt():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("hello.txt", "不是 Word 文件")
    assert _code(buf.getvalue(), "odd.docx") in (CORRUPT, UNSUPPORTED_FORMAT)


def test_zip_bomb_is_too_large():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"\x00" * 2_000_000)
    data = buf.getvalue()
    assert len(data) < 100_000
    assert _code(data, "bomb.docx", max_bytes=100_000, unzip_ratio=8) == TOO_LARGE


# ── pptx ────────────────────────────────────────────────────────────


def _deck(presentation):
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "季度報告：營運概況與下一季目標"
    first.placeholders[1].text = "營收成長百分之二十\n成本下降，毛利率提升到四成"
    first.notes_slide.notes_text_frame.text = "講者備註：強調成長來自新客戶"
    presentation.slides.add_slide(presentation.slide_layouts[6])  # 空白投影片
    third = presentation.slides.add_slide(presentation.slide_layouts[5])
    third.shapes.title.text = "表格與群組"
    table = third.shapes.add_table(
        2, 2, Inches(1), Inches(2), Inches(4), Inches(1)
    ).table
    for (row, col), text in {
        (0, 0): "項目",
        (0, 1): "數值",
        (1, 0): "用戶",
        (1, 1): "1200",
    }.items():
        table.cell(row, col).text = text
    group = third.shapes.add_group_shape()
    inner = group.shapes.add_group_shape()
    box = inner.shapes.add_textbox(Inches(1), Inches(4), Inches(2), Inches(1))
    box.text_frame.text = "巢狀群組內文字"


def test_pptx_slides_title_body_table_group_notes(make_pptx):
    result = extract(make_pptx(_deck), "報告.pptx")
    assert result.format == "pptx"
    assert _locators(result) == [
        {"kind": "slide", "value": 1},
        {"kind": "slide", "value": 3},  # 空白投影片不產 segment
    ]
    first, third = (s.text for s in result.segments)
    assert first.split("\n\n") == [
        "季度報告：營運概況與下一季目標",
        "營收成長百分之二十\n成本下降，毛利率提升到四成",
        "講者備註：強調成長來自新客戶",
    ]
    assert third.split("\n\n") == [
        "表格與群組",
        "項目 | 數值\n用戶 | 1200",
        "巢狀群組內文字",
    ]
    assert first.count("季度報告") == 1  # 標題不重複計入


def test_pptx_error_classification(make_pptx, make_docx):
    good = make_pptx(_deck)
    assert _code(good[: len(good) // 3], "broken.pptx") == CORRUPT
    docx_bytes = make_docx(lambda d: d.add_paragraph("內容" * 40))
    assert _code(docx_bytes, "renamed.pptx") == UNSUPPORTED_FORMAT
    encrypted = OLE_MAGIC + b"\x00" * 64 + "EncryptedPackage".encode("utf-16-le")
    assert _code(encrypted, "secret.pptx") == ENCRYPTED


def test_pptx_images_only_is_empty_extraction(make_pptx):
    def build(presentation):
        for _ in range(3):
            presentation.slides.add_slide(presentation.slide_layouts[6])

    assert _code(make_pptx(build), "photos.pptx") == EMPTY_EXTRACTION
