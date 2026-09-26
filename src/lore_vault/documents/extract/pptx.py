"""pptx 抽取器（T-62）：`python-pptx`（本模組內的 `import pptx` 是絕對匯入，
指向第三方套件，不是本檔）。

每張投影片一個 segment（locator `slide`，1 起算），依序包含：標題、其餘形狀的
文字（群組形狀遞迴走訪）、表格（每列一行、` | ` 分隔）、備註。沒有任何文字的
投影片不產 segment。
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from typing import Any

import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .base import (
    CORRUPT,
    UNSUPPORTED_FORMAT,
    Budget,
    ExtractionError,
    Limits,
    Locator,
    Segment,
    check_ooxml_container,
)


def _is_group(shape: Any) -> bool:
    try:
        return shape.shape_type == MSO_SHAPE_TYPE.GROUP
    except (NotImplementedError, ValueError, KeyError):
        return False


def _shape_texts(shape: Any) -> Iterator[str]:
    if _is_group(shape):
        for child in shape.shapes:
            yield from _shape_texts(child)
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text:
            yield text
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = []
        for row in shape.table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            yield "\n".join(rows)


def _slide_text(slide: Any) -> str:
    parts: list[str] = []
    title = slide.shapes.title
    title_id = None
    if title is not None:
        title_id = title.shape_id
        if title.has_text_frame and title.text_frame.text.strip():
            parts.append(title.text_frame.text.strip())
    for shape in slide.shapes:
        if title_id is not None and shape.shape_id == title_id:
            continue
        parts.extend(_shape_texts(shape))
    if slide.has_notes_slide:
        frame = slide.notes_slide.notes_text_frame
        if frame is not None and frame.text.strip():
            parts.append(frame.text.strip())
    return "\n\n".join(parts)


def extract_pptx(data: bytes, budget: Budget, limits: Limits) -> list[Segment]:
    check_ooxml_container(data, limits, label="pptx")
    try:
        presentation = pptx.Presentation(io.BytesIO(data))
    except ValueError as exc:
        # zip 合法但主文件內容型別不是簡報（例如 docx 改了副檔名）
        raise ExtractionError(UNSUPPORTED_FORMAT, f"不是 pptx：{exc}") from None
    except Exception as exc:
        raise ExtractionError(CORRUPT, f"pptx 解析失敗：{exc}") from None
    segments: list[Segment] = []
    try:
        for number, slide in enumerate(presentation.slides, start=1):
            text = _slide_text(slide)
            if text.strip():
                segments.append(Segment(budget.add(text), Locator("slide", number)))
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(CORRUPT, f"pptx 內容解析失敗：{exc}") from None
    return segments
